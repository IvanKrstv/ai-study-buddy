from fastapi import UploadFile, HTTPException

from zipfile import BadZipFile
from docx.opc.exceptions import PackageNotFoundError
from pypdf import PdfReader
from pypdf.errors import PdfReadError, FileNotDecryptedError
from docx import Document
import io

from ai_study_buddy.services.file_validator import FileValidator


class ExtractFile:
    def __init__(self, file: UploadFile, validator: FileValidator | None = None):
        self._file = file
        # composition validator
        self._validator = validator or FileValidator()


    async def extract_text(self) -> str:
        # Check if the file is too big
        self._validator.validate_file_size(self._file.size)

        # read the file content
        content = await self._file.read()
        file_extension = self._file.filename.split('.')[-1].lower()
        self._validator.validate_extension(file_extension)
        self._validator.validate_content_matches_extension(file_extension, content)

        match file_extension:
            case 'pdf':
                return self._extract_from_pdf(content)
            case 'docx':
                return self._extract_from_docx(content)
            case 'txt':
                return self._extract_from_txt(content)
            case _: # cannot be reached, cases are validated
                raise AssertionError('unreachable')


    def _extract_from_pdf(self, content: bytes) -> str:
        try:
            reader = PdfReader(stream=io.BytesIO(content))
            text = []
            for page in reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text.append(extracted)
        except (PdfReadError, FileNotDecryptedError):
            raise HTTPException(
                status_code=422,
                detail="The PDF file is corrupted or password protected."
            )

        self._validator.empty_text_content_check(text, error_message='The uploaded PDF contains only scanned images and no extractable text. OCR is required.')

        return '\n'.join(text)


    def _extract_from_docx(self, content: bytes) -> str:
        text = []

        try:
            doc = Document(io.BytesIO(content))
            for paragraph in doc.paragraphs:
                extracted_text = paragraph.text
                if extracted_text:
                    text.append(extracted_text)
            for table in doc.tables:
                for row in table.rows:
                    row_text = ' | '.join(cell.text for cell in row.cells)
                    if row_text.strip(' |'):
                        text.append(row_text)
        except (PackageNotFoundError, BadZipFile):
            raise HTTPException(
                status_code=422,
                detail="The docx file is corrupted."
            )

        self._validator.empty_text_content_check(text, error_message='The document is empty.')

        return '\n'.join(text)


    def _extract_from_txt(self, content: bytes) -> str:
        try:
            text = content.decode('utf-8')
        except UnicodeDecodeError:
            raise HTTPException(
                status_code=422,
                detail="Could not read text file — please make sure it's UTF-8 encoded."
            )

        self._validator.empty_text_content_check(text, error_message='The document is empty.')

        return text
